import uuid
import io
import pytest
from unittest.mock import patch

@pytest.fixture(autouse=True)
def mock_embedding_and_vector_services():
    with patch("app.services.ingestion_pipeline.embed_texts", return_value=[[0.1, 0.2, 0.3]]), \
         patch("app.services.ingestion_pipeline.index_chunks", return_value=None):
        yield

def test_create_document_metadata(client):
    response = client.post(
        "/api/documents",
        json={"title": "Test Doc", "source_type": "pdf", "original_filename": "test.pdf"}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["title"] == "Test Doc"
    assert data["source_type"] == "pdf"
    assert data["original_filename"] == "test.pdf"
    assert data["status"] == "pending"
    assert "id" in data
    assert "created_at" in data
    assert "updated_at" in data
    
    # Check UUID format
    assert uuid.UUID(data["id"])

def test_list_documents(client):
    client.post("/api/documents", json={"title": "Doc 1", "source_type": "text"})
    response = client.get("/api/documents")
    assert response.status_code == 200
    data = response.json()
    assert data["total"] == 1
    assert len(data["items"]) == 1

def test_get_document(client):
    create_resp = client.post("/api/documents", json={"title": "Doc 1", "source_type": "text"})
    doc_id = create_resp.json()["id"]
    
    response = client.get(f"/api/documents/{doc_id}")
    assert response.status_code == 200
    data = response.json()
    assert data["id"] == doc_id
    assert "created_at" in data
    assert "updated_at" in data

def test_update_document(client):
    create_resp = client.post("/api/documents", json={"title": "Doc 1", "source_type": "text"})
    doc_id = create_resp.json()["id"]
    
    response = client.patch(f"/api/documents/{doc_id}", json={"status": "processing"})
    assert response.status_code == 200
    assert response.json()["status"] == "processing"

def test_delete_document(client):
    create_resp = client.post("/api/documents", json={"title": "Doc 1", "source_type": "text"})
    doc_id = create_resp.json()["id"]
    
    response = client.delete(f"/api/documents/{doc_id}")
    assert response.status_code == 204
    
    get_resp = client.get(f"/api/documents/{doc_id}")
    assert get_resp.status_code == 404

def test_create_duplicate_collection(client):
    response = client.post("/api/collections", json={"name": "Science"})
    assert response.status_code == 200
    
    response2 = client.post("/api/collections", json={"name": "Science"})
    assert response2.status_code == 409

def test_create_duplicate_tag(client):
    response = client.post("/api/tags", json={"name": "important"})
    assert response.status_code == 200
    
    response2 = client.post("/api/tags", json={"name": "important"})
    assert response2.status_code == 409

def test_attach_document_to_collection(client):
    doc_resp = client.post("/api/documents", json={"title": "Doc", "source_type": "pdf"})
    doc_id = doc_resp.json()["id"]
    
    col_resp = client.post("/api/collections", json={"name": "My Collection"})
    col_id = col_resp.json()["id"]
    
    attach_resp = client.post(f"/api/collections/{col_id}/documents/{doc_id}")
    assert attach_resp.status_code == 200
    data = attach_resp.json()
    assert len(data["collections"]) == 1
    assert data["collections"][0]["id"] == col_id

def test_attach_tag_to_document(client):
    doc_resp = client.post("/api/documents", json={"title": "Doc", "source_type": "pdf"})
    doc_id = doc_resp.json()["id"]
    
    tag_resp = client.post("/api/tags", json={"name": "My Tag"})
    tag_id = tag_resp.json()["id"]
    
    attach_resp = client.post(f"/api/tags/{tag_id}/documents/{doc_id}")
    assert attach_resp.status_code == 200
    data = attach_resp.json()
    assert len(data["tags"]) == 1
    assert data["tags"][0]["id"] == tag_id

    detach_resp = client.delete(f"/api/tags/{tag_id}/documents/{doc_id}")
    assert detach_resp.status_code == 200
    assert detach_resp.json()["tags"] == []

def test_document_jobs_endpoint(client, db):
    # Setup document and job directly to test endpoint
    doc_resp = client.post("/api/documents", json={"title": "Doc", "source_type": "pdf"})
    doc_id = doc_resp.json()["id"]
    
    from app.models.document import IngestionJob
    job = IngestionJob(user_id="dev-user-id", document_id=doc_id, job_type="parse", status="queued")
    db.add(job)
    db.commit()
    
    jobs_resp = client.get(f"/api/documents/{doc_id}/jobs")
    assert jobs_resp.status_code == 200
    jobs = jobs_resp.json()
    assert len(jobs) == 1
    assert jobs[0]["job_type"] == "parse"

def test_invalid_source_type(client):
    response = client.post(
        "/api/documents",
        json={"title": "Test", "source_type": "invalid_type"}
    )
    assert response.status_code == 422

def test_upload_txt_file(client, db):
    file_content = b"This is a test document.\n\nIt has multiple paragraphs.\n\n" * 10
    files = {"file": ("test.txt", io.BytesIO(file_content), "text/plain")}
    data = {"title": "Test TXT Upload"}
    
    response = client.post("/api/documents/upload", files=files, data=data)
    assert response.status_code == 200
    resp_data = response.json()
    assert resp_data["document"]["title"] == "Test TXT Upload"
    assert resp_data["document"]["source_type"] == "text"
    assert resp_data["document"]["status"] == "processing"
    assert resp_data["job"]["status"] == "queued"
    assert resp_data["chunks_created"] == 0
    assert resp_data["message"] == "Ingestion queued"
    
    job_id = resp_data["job"]["id"]
    doc_id = resp_data["document"]["id"]
    
    # Run the background task directly
    from app.services.job_runner import run_ingestion_job
    import os
    from app.models.document import Document, IngestionJob
    from app.core.config import settings
    file_path = os.path.join(settings.STORAGE_DIR, "originals", doc_id, "test.txt")
    
    run_ingestion_job(job_id, "file_upload", {"file_path": file_path, "source_type": "text"}, db=db)
    
    # Check status endpoint
    status_resp = client.get(f"/api/jobs/{job_id}")
    assert status_resp.status_code == 200
    status_data = status_resp.json()
    assert status_data["job"]["status"] == "succeeded"
    assert status_data["document"]["status"] == "ready"
    assert status_data["chunks_created"] > 0
    
    chunks_resp = client.get(f"/api/documents/{doc_id}/chunks")
    assert chunks_resp.status_code == 200
    chunks_data = chunks_resp.json()
    assert chunks_data["total"] == status_data["chunks_created"]
    assert chunks_data["items"][0]["chunk_index"] == 0

def test_upload_markdown_file(client, db):
    file_content = b"# Header\n\nSome markdown text.\n"
    files = {"file": ("test.md", io.BytesIO(file_content), "text/markdown")}
    
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 200
    data = response.json()
    assert data["document"]["source_type"] == "markdown"
    assert data["job"]["status"] == "queued"
    
    from app.services.job_runner import run_ingestion_job
    import os
    from app.models.document import Document, IngestionJob
    from app.core.config import settings
    file_path = os.path.join(settings.STORAGE_DIR, "originals", data["document"]["id"], "test.md")
    run_ingestion_job(data["job"]["id"], "file_upload", {"file_path": file_path, "source_type": "markdown"}, db=db)
    
    doc = db.query(Document).filter(Document.id == data["document"]["id"]).first()
    assert doc.status == "ready"

def test_upload_unsupported_file(client):
    file_content = b"fake image"
    files = {"file": ("test.png", io.BytesIO(file_content), "image/png")}
    
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 415

def test_upload_empty_txt_file(client, db):
    file_content = b"   \n  "
    files = {"file": ("empty.txt", io.BytesIO(file_content), "text/plain")}
    
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 200
    data = response.json()
    assert data["document"]["status"] == "processing"
    
    from app.services.job_runner import run_ingestion_job
    import os
    from app.models.document import Document, IngestionJob
    from app.core.config import settings
    file_path = os.path.join(settings.STORAGE_DIR, "originals", data["document"]["id"], "empty.txt")
    run_ingestion_job(data["job"]["id"], "file_upload", {"file_path": file_path, "source_type": "text"}, db=db)
    
    doc = db.query(Document).filter(Document.id == data["document"]["id"]).first()
    job = db.query(IngestionJob).filter(IngestionJob.id == data["job"]["id"]).first()
    assert doc.status == "failed"
    assert job.status == "failed"
    assert "empty" in job.error_message.lower()

def test_get_chunks_missing_doc(client):
    response = client.get("/api/documents/00000000-0000-0000-0000-000000000000/chunks")
    assert response.status_code == 404

def test_delete_document_cascades_chunks(client, db):
    file_content = b"This is a test document."
    files = {"file": ("test.txt", io.BytesIO(file_content), "text/plain")}
    
    upload_resp = client.post("/api/documents/upload", files=files)
    data = upload_resp.json()
    doc_id = data["document"]["id"]
    
    from app.services.job_runner import run_ingestion_job
    import os
    from app.models.document import Document, IngestionJob
    from app.core.config import settings
    file_path = os.path.join(settings.STORAGE_DIR, "originals", doc_id, "test.txt")
    run_ingestion_job(data["job"]["id"], "file_upload", {"file_path": file_path, "source_type": "text"}, db=db)
    
    chunks_resp = client.get(f"/api/documents/{doc_id}/chunks")
    assert chunks_resp.json()["total"] > 0
    
    del_resp = client.delete(f"/api/documents/{doc_id}")
    assert del_resp.status_code == 204
    
    doc_resp = client.get(f"/api/documents/{doc_id}")
    assert doc_resp.status_code == 404

def test_upload_docx_file(client, db):
    from docx import Document as DocxDocument
    import io
    
    doc = DocxDocument()
    doc.add_paragraph('This is a test DOCX document.')
    
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    
    files = {"file": ("test.docx", doc_io, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 200
    resp_data = response.json()
    assert resp_data["document"]["source_type"] == "docx"
    assert resp_data["document"]["status"] == "processing"
    
    from app.services.job_runner import run_ingestion_job
    import os
    from app.models.document import Document, IngestionJob
    from app.core.config import settings
    file_path = os.path.join(settings.STORAGE_DIR, "originals", resp_data["document"]["id"], "test.docx")
    run_ingestion_job(resp_data["job"]["id"], "file_upload", {"file_path": file_path, "source_type": "docx"}, db=db)
    
    doc = db.query(Document).filter(Document.id == resp_data["document"]["id"]).first()
    assert doc.status == "ready"
    
    
    doc_id = resp_data["document"]["id"]
    chunks_resp = client.get(f"/api/documents/{doc_id}/chunks")
    assert chunks_resp.json()["total"] > 0

# Note on PDF Testing: 
# PDF extraction is fully supported via pypdf but generating a valid PDF purely in memory 
# without heavy external dependencies (like reportlab) is overly complex for unit testing. 
# It is excluded from the test suite but functionally present in extraction_service.py.

def test_index_document_not_found(client):
    response = client.post("/api/documents/00000000-0000-0000-0000-000000000000/index")
    assert response.status_code == 404

def test_index_document_no_chunks(client):
    doc_resp = client.post("/api/documents", json={"title": "Doc", "source_type": "text"})
    doc_id = doc_resp.json()["id"]
    
    response = client.post(f"/api/documents/{doc_id}/index")
    assert response.status_code == 400
    assert "no chunks" in response.json()["detail"]

def test_index_document_success(client, db):
    file_content = b"This is a test document."
    files = {"file": ("test.txt", io.BytesIO(file_content), "text/plain")}
    
    upload_resp = client.post("/api/documents/upload", files=files)
    data = upload_resp.json()
    doc_id = data["document"]["id"]
    
    from app.services.job_runner import run_ingestion_job
    import os
    from app.models.document import Document, IngestionJob
    from app.core.config import settings
    file_path = os.path.join(settings.STORAGE_DIR, "originals", doc_id, "test.txt")
    run_ingestion_job(data["job"]["id"], "file_upload", {"file_path": file_path, "source_type": "text"}, db=db)
    
    # upload already indexes, let's re-index
    response = client.post(f"/api/documents/{doc_id}/index")
    assert response.status_code == 200
    
def test_upload_indexing_failure_fallback(client, db):
    with patch("app.services.job_runner.index_document_chunks", return_value={"failed_count": 1, "indexed_count": 0, "message": "Chroma down"}):
        file_content = b"This is a test document."
        files = {"file": ("test_fail.txt", io.BytesIO(file_content), "text/plain")}
        
    assert response.status_code == 422

def test_upload_txt_file(client, db):
    file_content = b"This is a test document.\n\nIt has multiple paragraphs.\n\n" * 10
    files = {"file": ("test.txt", io.BytesIO(file_content), "text/plain")}
    data = {"title": "Test TXT Upload"}
    
    response = client.post("/api/documents/upload", files=files, data=data)
    assert response.status_code == 200
    resp_data = response.json()
    assert resp_data["document"]["title"] == "Test TXT Upload"
    assert resp_data["document"]["source_type"] == "text"
    assert resp_data["document"]["status"] == "processing"
    assert resp_data["job"]["status"] == "queued"
    assert resp_data["chunks_created"] == 0
    assert resp_data["message"] == "Ingestion queued"
    
    job_id = resp_data["job"]["id"]
    doc_id = resp_data["document"]["id"]
    
    # Run the background task directly
    from app.services.job_runner import run_ingestion_job
    import os
    from app.models.document import Document, IngestionJob
    from app.core.config import settings
    file_path = os.path.join(settings.STORAGE_DIR, "originals", doc_id, "test.txt")
    
    run_ingestion_job(job_id, "file_upload", {"file_path": file_path, "source_type": "text"}, db=db)
    
    # Check status endpoint
    status_resp = client.get(f"/api/jobs/{job_id}")
    assert status_resp.status_code == 200
    status_data = status_resp.json()
    assert status_data["job"]["status"] == "succeeded"
    assert status_data["document"]["status"] == "ready"
    assert status_data["chunks_created"] > 0
    
    chunks_resp = client.get(f"/api/documents/{doc_id}/chunks")
    assert chunks_resp.status_code == 200
    chunks_data = chunks_resp.json()
    assert chunks_data["total"] == status_data["chunks_created"]
    assert chunks_data["items"][0]["chunk_index"] == 0

def test_upload_markdown_file(client, db):
    file_content = b"# Header\n\nSome markdown text.\n"
    files = {"file": ("test.md", io.BytesIO(file_content), "text/markdown")}
    
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 200
    data = response.json()
    assert data["document"]["source_type"] == "markdown"
    assert data["job"]["status"] == "queued"
    
    from app.services.job_runner import run_ingestion_job
    import os
    from app.models.document import Document, IngestionJob
    from app.core.config import settings
    file_path = os.path.join(settings.STORAGE_DIR, "originals", data["document"]["id"], "test.md")
    run_ingestion_job(data["job"]["id"], "file_upload", {"file_path": file_path, "source_type": "markdown"}, db=db)
    
    doc = db.query(Document).filter(Document.id == data["document"]["id"]).first()
    assert doc.status == "ready"

def test_upload_unsupported_file(client):
    file_content = b"fake image"
    files = {"file": ("test.png", io.BytesIO(file_content), "image/png")}
    
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 415

def test_upload_empty_txt_file(client, db):
    file_content = b"   \n  "
    files = {"file": ("empty.txt", io.BytesIO(file_content), "text/plain")}
    
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 200
    data = response.json()
    assert data["document"]["status"] == "processing"
    
    from app.services.job_runner import run_ingestion_job
    import os
    from app.models.document import Document, IngestionJob
    from app.core.config import settings
    file_path = os.path.join(settings.STORAGE_DIR, "originals", data["document"]["id"], "empty.txt")
    run_ingestion_job(data["job"]["id"], "file_upload", {"file_path": file_path, "source_type": "text"}, db=db)
    
    doc = db.query(Document).filter(Document.id == data["document"]["id"]).first()
    job = db.query(IngestionJob).filter(IngestionJob.id == data["job"]["id"]).first()
    assert doc.status == "failed"
    assert job.status == "failed"
    assert "empty" in job.error_message.lower()

def test_get_chunks_missing_doc(client):
    response = client.get("/api/documents/00000000-0000-0000-0000-000000000000/chunks")
    assert response.status_code == 404

def test_delete_document_cascades_chunks(client, db):
    file_content = b"This is a test document."
    files = {"file": ("test.txt", io.BytesIO(file_content), "text/plain")}
    
    upload_resp = client.post("/api/documents/upload", files=files)
    data = upload_resp.json()
    doc_id = data["document"]["id"]
    
    from app.services.job_runner import run_ingestion_job
    import os
    from app.models.document import Document, IngestionJob
    from app.core.config import settings
    file_path = os.path.join(settings.STORAGE_DIR, "originals", doc_id, "test.txt")
    run_ingestion_job(data["job"]["id"], "file_upload", {"file_path": file_path, "source_type": "text"}, db=db)
    
    chunks_resp = client.get(f"/api/documents/{doc_id}/chunks")
    assert chunks_resp.json()["total"] > 0
    
    del_resp = client.delete(f"/api/documents/{doc_id}")
    assert del_resp.status_code == 204
    
    doc_resp = client.get(f"/api/documents/{doc_id}")
    assert doc_resp.status_code == 404

def test_upload_docx_file(client, db):
    from docx import Document as DocxDocument
    import io
    
    doc = DocxDocument()
    doc.add_paragraph('This is a test DOCX document.')
    
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    
    files = {"file": ("test.docx", doc_io, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 200
    resp_data = response.json()
    assert resp_data["document"]["source_type"] == "docx"
    assert resp_data["document"]["status"] == "processing"
    
    from app.services.job_runner import run_ingestion_job
    import os
    from app.models.document import Document, IngestionJob
    from app.core.config import settings
    file_path = os.path.join(settings.STORAGE_DIR, "originals", resp_data["document"]["id"], "test.docx")
    run_ingestion_job(resp_data["job"]["id"], "file_upload", {"file_path": file_path, "source_type": "docx"}, db=db)
    
    doc = db.query(Document).filter(Document.id == resp_data["document"]["id"]).first()
    assert doc.status == "ready"
    
    
    doc_id = resp_data["document"]["id"]
    chunks_resp = client.get(f"/api/documents/{doc_id}/chunks")
    assert chunks_resp.json()["total"] > 0

# Note on PDF Testing: 
# PDF extraction is fully supported via pypdf but generating a valid PDF purely in memory 
# without heavy external dependencies (like reportlab) is overly complex for unit testing. 
# It is excluded from the test suite but functionally present in extraction_service.py.

def test_index_document_not_found(client):
    response = client.post("/api/documents/00000000-0000-0000-0000-000000000000/index")
    assert response.status_code == 404

def test_index_document_no_chunks(client):
    doc_resp = client.post("/api/documents", json={"title": "Doc", "source_type": "text"})
    doc_id = doc_resp.json()["id"]
    
    response = client.post(f"/api/documents/{doc_id}/index")
    assert response.status_code == 400
    assert "no chunks" in response.json()["detail"]

def test_index_document_success(client, db):
    file_content = b"This is a test document."
    files = {"file": ("test.txt", io.BytesIO(file_content), "text/plain")}
    
    upload_resp = client.post("/api/documents/upload", files=files)
    data = upload_resp.json()
    doc_id = data["document"]["id"]
    
    from app.services.job_runner import run_ingestion_job
    import os
    from app.models.document import Document, IngestionJob
    from app.core.config import settings
    file_path = os.path.join(settings.STORAGE_DIR, "originals", doc_id, "test.txt")
    run_ingestion_job(data["job"]["id"], "file_upload", {"file_path": file_path, "source_type": "text"}, db=db)
    
    # upload already indexes, let's re-index
    response = client.post(f"/api/documents/{doc_id}/index")
    assert response.status_code == 200
    
def test_upload_indexing_failure_fallback(client, db):
    with patch("app.services.job_runner.index_document_chunks", return_value={"failed_count": 1, "indexed_count": 0, "message": "Chroma down"}):
        file_content = b"This is a test document."
        files = {"file": ("test_fail.txt", io.BytesIO(file_content), "text/plain")}
        
        response = client.post("/api/documents/upload", files=files)
        assert response.status_code == 200 # Upload still succeeds
        data = response.json()
        assert data["document"]["status"] == "processing"
        
        from app.services.job_runner import run_ingestion_job
        import os
        from app.models.document import Document, IngestionJob
        from app.core.config import settings
        file_path = os.path.join(settings.STORAGE_DIR, "originals", data["document"]["id"], "test_fail.txt")
        run_ingestion_job(data["job"]["id"], "file_upload", {"file_path": file_path, "source_type": "text"}, db=db)
        
        doc = db.query(Document).filter(Document.id == data["document"]["id"]).first()
        job = db.query(IngestionJob).filter(IngestionJob.id == data["job"]["id"]).first()
        
        assert doc.status == "ready"
        assert job.status == "succeeded"
        assert "indexing failed" in job.message.lower()
        
        doc_id = data["document"]["id"]
        chunks_resp = client.get(f"/api/documents/{doc_id}/chunks")
        assert chunks_resp.status_code == 200
